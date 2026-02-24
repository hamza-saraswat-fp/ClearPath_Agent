"""ClearPath Import Template Parser.

Parses ClearPath import-format Excel files (the curated workflow templates).
These have a fixed 3-sheet structure with a horizontal status layout,
which differs from the internal/CSV format handled by ProductionDataParser.

Outputs the same ParsedWorkflow models so it plugs directly into the
existing ingestion pipeline.
"""

import logging
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from ..models.enums import StatusCategory, UserRole
from .prod_data_parser import (
    ActionButtonConfig,
    FocusViewConfig,
    ParsedWorkflow,
    StatusConfig,
)

logger = logging.getLogger(__name__)

# ClearPath import status type → StatusCategory
_STATUS_TYPE_MAP = {
    "new": StatusCategory.PENDING,
    "any": StatusCategory.IN_PROGRESS,
    "travel time": StatusCategory.IN_PROGRESS,
    "in progress": StatusCategory.IN_PROGRESS,
    "pending": StatusCategory.ON_HOLD,
    "complete": StatusCategory.COMPLETED,
    "canceled": StatusCategory.CANCELLED,
}

# ClearPath import user role strings → UserRole
_USER_ROLE_MAP = {
    "service agent": UserRole.SERVICE_AGENT,
    "technician": UserRole.TECHNICIAN,
    "tech": UserRole.TECHNICIAN,
    "office staff": UserRole.OFFICE_STAFF,
    "office": UserRole.OFFICE_STAFF,
    "manager": UserRole.MANAGER,
    "admin": UserRole.ADMIN,
    "dispatcher": UserRole.DISPATCHER,
    "sales": UserRole.SALES,
}


class TemplateParser:
    """Parser for ClearPath import-format Excel template files."""

    @staticmethod
    def is_template_format(filepath: Path) -> bool:
        """Check if an Excel file uses the ClearPath import template format.

        The signature is "Status Name 1" appearing in the first sheet's header row.
        """
        try:
            wb = load_workbook(filepath, read_only=True, data_only=True)
            first_sheet = wb[wb.sheetnames[0]]
            headers = [
                str(c.value).strip().lower() if c.value else ""
                for c in list(first_sheet.iter_rows(min_row=1, max_row=1))[0]
            ]
            wb.close()
            return any(h == "status name 1" for h in headers)
        except Exception:
            return False

    def parse_template(self, filepath: Path) -> ParsedWorkflow:
        """Parse a ClearPath import template Excel file.

        Args:
            filepath: Path to the template Excel file

        Returns:
            ParsedWorkflow with statuses, action buttons, and focus view configs
        """
        logger.info(f"Parsing template: {filepath.name}")

        wb = load_workbook(filepath, read_only=True, data_only=True)
        sheets = {name.lower(): wb[name] for name in wb.sheetnames}

        parsed = ParsedWorkflow(source_file=filepath.name)

        # Parse each sheet by name prefix
        for sheet_key, sheet in sheets.items():
            if "job custom status" in sheet_key:
                parsed.statuses = self._parse_statuses(sheet)
            elif "action button" in sheet_key:
                parsed.action_buttons = self._parse_action_buttons(sheet)
            elif "focus view" in sheet_key:
                parsed.focus_views = self._parse_focus_views(sheet)
            else:
                logger.warning(f"Skipping unknown template sheet: {sheet_key}")

        # Set flow name from statuses (col 0 of status sheet)
        if parsed.statuses:
            parsed.flow_name = parsed.statuses[0].flow_name

        wb.close()

        logger.info(
            f"Parsed template {filepath.name}: "
            f"{parsed.status_count} statuses, "
            f"{parsed.action_count} actions, "
            f"{parsed.widget_count} unique widgets"
        )

        return parsed

    def _parse_statuses(self, sheet: Worksheet) -> list[StatusConfig]:
        """Parse horizontal status layout.

        Col 0: Workflow Name
        Cols 1+: repeating groups of 4 (Status Name N, Type N, Color N, Icon N)
        """
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2:
            return []

        data_row = rows[1]
        flow_name = _clean(data_row[0]) if data_row[0] else ""

        statuses = []
        col = 1
        seq = 1

        while col + 3 < len(data_row):
            name = _clean(data_row[col]) if data_row[col] else None
            if not name:
                break

            type_str = _clean(data_row[col + 1]).lower() if data_row[col + 1] else ""
            color = _clean(data_row[col + 2]) if data_row[col + 2] else "#3B82F6"
            if not color.startswith("#"):
                color = "#3B82F6"

            category = _STATUS_TYPE_MAP.get(type_str, StatusCategory.IN_PROGRESS)

            statuses.append(StatusConfig(
                name=name,
                category=category,
                sequence=seq,
                color=color,
                flow_name=flow_name,
                is_active=True,
            ))

            col += 4
            seq += 1

        return statuses

    def _parse_action_buttons(self, sheet: Worksheet) -> list[ActionButtonConfig]:
        """Parse action buttons (vertical layout, 7 fixed columns).

        Col 0: Custom Job Status Workflow Name
        Col 1: Status Action Flow Name
        Col 2: Job Status Name
        Col 3: User Role
        Col 4: Button Action
        Col 5: Action Option (target status for Change Status)
        Col 6: Action Button Name
        """
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2:
            return []

        buttons = []

        for row_num, row in enumerate(rows[1:], start=2):
            if not row or not any(row):
                continue

            # Need at least 5 columns (through Button Action)
            if len(row) < 5:
                continue

            action_type = _clean(row[4]) if row[4] else None
            if not action_type:
                continue

            flow_name = _clean(row[0]) if row[0] else ""
            status_name = _clean(row[2]) if row[2] else ""
            user_role = _parse_role(row[3]) if row[3] else UserRole.SERVICE_AGENT
            action_option = _clean(row[5]) if len(row) > 5 and row[5] else None
            label = _clean(row[6]) if len(row) > 6 and row[6] else action_type

            buttons.append(ActionButtonConfig(
                action_type=action_type,
                label=label,
                status_name=status_name,
                flow_name=flow_name,
                user_role=user_role,
                order=row_num - 1,
                template_id=action_option,  # target status for Change Status
            ))

        return buttons

    def _parse_focus_views(self, sheet: Worksheet) -> list[FocusViewConfig]:
        """Parse focus view + status instruction (vertical layout, 10 fixed columns).

        Col 0: Custom Job Status Workflow Name
        Col 1: Status Action Flow Name
        Col 2: Job Status Name
        Col 3: User Role
        Col 4: Status Instructions
        Col 5: Display Action Menu (T/F)
        Col 6: Ability to Change Status (T/F)
        Col 7: Focus View Enabled (T/F)
        Col 8: Focus View Layout (comma-separated widgets)
        Col 9: Restrict user access to Focus View only (T/F)
        """
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2:
            return []

        focus_views = []

        for row_num, row in enumerate(rows[1:], start=2):
            if not row or not any(row):
                continue

            # Need at least 3 columns (through Job Status Name)
            if len(row) < 3:
                continue

            status_name = _clean(row[2]) if row[2] else None
            if not status_name:
                continue

            flow_name = _clean(row[0]) if row[0] else ""
            user_role = _parse_role(row[3]) if len(row) > 3 and row[3] else UserRole.SERVICE_AGENT
            instructions = _clean(row[4]) if len(row) > 4 and row[4] else ""

            # Parse booleans (T/F strings)
            display_action_menu = _parse_tf(row[5]) if len(row) > 5 else True
            ability_to_change = _parse_tf(row[6]) if len(row) > 6 else True
            fv_enabled = _parse_tf(row[7]) if len(row) > 7 else True
            restrict = _parse_tf(row[9]) if len(row) > 9 else False

            # Parse widgets from comma-separated Focus View Layout
            widgets = []
            if len(row) > 8 and row[8]:
                widget_str = _clean(row[8])
                widgets = [w.strip() for w in widget_str.split(",") if w.strip()]

            focus_views.append(FocusViewConfig(
                status_name=status_name,
                flow_name=flow_name,
                user_role=user_role,
                widgets=widgets,
                status_instructions=instructions,
                display_action_menu=display_action_menu,
                ability_to_change_status=ability_to_change,
                focus_view_enabled=fv_enabled,
                restrict_to_focus_view=restrict,
            ))

        return focus_views


def _clean(value) -> str:
    """Strip whitespace from a cell value."""
    return str(value).strip() if value is not None else ""


def _parse_tf(value) -> bool:
    """Parse T/F boolean strings."""
    if value is None:
        return False
    return str(value).strip().upper() in ("T", "TRUE", "YES", "1")


def _parse_role(value) -> UserRole:
    """Parse user role string."""
    if value is None:
        return UserRole.SERVICE_AGENT
    return _USER_ROLE_MAP.get(str(value).strip().lower(), UserRole.SERVICE_AGENT)
