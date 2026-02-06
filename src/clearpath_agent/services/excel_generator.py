"""Excel Generator Service.

Generates physical .xlsx files from ExcelImportTemplate schemas
using openpyxl. Matches FieldPulse ClearPath Import Template format exactly.
"""

import logging
from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from ..models.excel_schemas import ExcelImportTemplate

logger = logging.getLogger(__name__)


class ExcelGenerator:
    """Generates Excel files matching FieldPulse ClearPath Import Template format.

    Creates properly formatted .xlsx files with 3 tabs:
    1. Job Custom Status - HORIZONTAL layout (all statuses in one row)
    2. Action Buttons - Correct column names
    3. Focus View + Status Instruction - Correct column names
    """

    # Sheet names matching FieldPulse import format
    SHEET_JOB_CUSTOM_STATUS = "Job Custom Status"
    SHEET_ACTION_BUTTONS = "Action Buttons"
    SHEET_FOCUS_VIEW = "Focus View + Status Instruction"

    # Column headers for Action Buttons tab (matches FieldPulse exactly)
    HEADERS_ACTION_BUTTONS = [
        "Custom Job Status Workflow Name",
        "Status Action Flow Name",
        "Job Status Name",
        "User Role",
        "Button Action",
        "Action Option",
        "Action Button Name",
    ]

    # Column headers for Focus View tab (matches FieldPulse exactly)
    HEADERS_FOCUS_VIEW = [
        "Custom Job Status Workflow Name",
        "Status Action Flow Name",
        "Job Status Name",
        "User Role",
        "Status Instructions",
        "Display Action Menu",
        "Ability to Change Status",
        "Focus View Enabled",
        "Focus View Layout",
        "Restrict user access to Focus View only",
    ]

    # Column widths
    WIDTHS_ACTION_BUTTONS = [35, 30, 20, 15, 25, 30, 25]
    WIDTHS_FOCUS_VIEW = [35, 30, 20, 15, 40, 18, 22, 18, 50, 30]

    # Styling constants
    HEADER_FONT = Font(bold=True, size=11)
    HEADER_FILL = PatternFill(
        start_color="D3D3D3", end_color="D3D3D3", fill_type="solid"
    )
    DATA_ALIGNMENT = Alignment(horizontal="left", vertical="top")
    WRAP_ALIGNMENT = Alignment(horizontal="left", vertical="top", wrap_text=True)
    THIN_BORDER = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    def __init__(
        self,
        default_output_dir: Optional[Path] = None,
    ):
        """Initialize the Excel generator.

        Args:
            default_output_dir: Default directory for output files.
                Defaults to 'reports/' in the project root.
        """
        self.default_output_dir = default_output_dir or Path("reports")

    def generate(
        self,
        template: ExcelImportTemplate,
        output_path: Optional[Path] = None,
        overwrite: bool = False,
        validate_first: bool = True,
    ) -> Path:
        """Generate an Excel file from the template.

        Args:
            template: The ExcelImportTemplate to convert to Excel
            output_path: Optional path for the output file.
                If not provided, generates a timestamped filename.
            overwrite: Whether to overwrite existing files.
                Raises error if False and file exists.
            validate_first: Whether to validate the template before generation.

        Returns:
            Path to the generated Excel file.

        Raises:
            ValueError: If template is invalid or file exists and overwrite=False.
            IOError: If file cannot be written.
        """
        # Validate template if requested
        if validate_first:
            errors = self._validate_template(template)
            if errors:
                raise ValueError(
                    f"Template validation failed:\n" + "\n".join(f"  - {e}" for e in errors)
                )

        # Determine output path
        output_path = self._resolve_output_path(template, output_path, overwrite)

        logger.info(f"Generating Excel file: {output_path}")

        # Create workbook and sheets
        workbook = Workbook()

        # Remove default sheet
        default_sheet = workbook.active
        if default_sheet is not None:
            workbook.remove(default_sheet)

        # Generate all three tabs
        self._generate_job_custom_status_sheet(workbook, template)
        self._generate_action_buttons_sheet(workbook, template)
        self._generate_focus_view_sheet(workbook, template)

        # Ensure output directory exists
        output_path.parent.mkdir(parents=True, exist_ok=True)

        # Save workbook
        try:
            workbook.save(output_path)
        except Exception as e:
            raise IOError(f"Failed to save Excel file: {e}") from e

        logger.info(
            f"Excel file generated successfully: {output_path} "
            f"({len(template.job_custom_statuses)} workflows, "
            f"{len(template.action_buttons)} button rows, "
            f"{len(template.focus_view)} focus view rows)"
        )

        return output_path.absolute()

    def _validate_template(self, template: ExcelImportTemplate) -> list[str]:
        """Validate the template before generation.

        Args:
            template: The template to validate

        Returns:
            List of validation error messages
        """
        errors = []

        # Check for empty template
        if not template.job_custom_statuses:
            errors.append("Template has no workflows defined")

        # Get all status names from all workflows
        status_names: set[str] = set()
        for workflow_row in template.job_custom_statuses:
            for status in workflow_row.statuses:
                status_names.add(status.name)

        # Check button rows reference valid statuses
        for row in template.action_buttons:
            if row.job_status_name not in status_names:
                errors.append(
                    f"Action button references unknown status: {row.job_status_name}"
                )

        # Check focus view rows reference valid statuses
        for row in template.focus_view:
            if row.job_status_name not in status_names:
                errors.append(
                    f"Focus view references unknown status: {row.job_status_name}"
                )

        return errors

    def _resolve_output_path(
        self,
        template: ExcelImportTemplate,
        output_path: Optional[Path],
        overwrite: bool,
    ) -> Path:
        """Resolve the output file path.

        Args:
            template: The template (used for default filename)
            output_path: Optional explicit output path
            overwrite: Whether to allow overwriting

        Returns:
            Resolved Path object

        Raises:
            ValueError: If file exists and overwrite=False
        """
        if output_path is None:
            # Generate default filename
            workflow_name = template.job_custom_statuses[0].workflow_name
            # Sanitize workflow name for filename
            safe_name = "".join(c if c.isalnum() or c in "-_ " else "_" for c in workflow_name)
            safe_name = safe_name.replace(" ", "_")
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"ClearPath_Import_{safe_name}_{timestamp}.xlsx"
            output_path = self.default_output_dir / filename
        else:
            output_path = Path(output_path)

        # Check for existing file
        if output_path.exists() and not overwrite:
            raise ValueError(
                f"Output file already exists: {output_path}. "
                "Set overwrite=True to replace."
            )

        return output_path

    def _generate_job_custom_status_sheet(
        self,
        workbook: Workbook,
        template: ExcelImportTemplate,
    ) -> None:
        """Generate the Job Custom Status tab with HORIZONTAL layout.

        FieldPulse format: One row per workflow with columns:
        Custom Job Status Workflow Name | Status Name 1 | Status Type 1 | Status Color 1 | Status Icon 1 | ...

        Args:
            workbook: The workbook to add the sheet to
            template: The template containing status data
        """
        sheet = workbook.create_sheet(self.SHEET_JOB_CUSTOM_STATUS)

        # Build dynamic headers based on max number of statuses
        max_statuses = template.get_max_statuses()
        headers = ["Custom Job Status Workflow Name"]
        for i in range(1, max_statuses + 1):
            headers.extend([
                f"Status Name {i}",
                f"Status Type {i}",
                f"Status Color {i}",
                f"Status Icon {i}",
            ])

        # Write header row
        self._write_header_row(sheet, headers)

        # Write data rows (one row per workflow)
        for row_idx, workflow_row in enumerate(template.job_custom_statuses, start=2):
            data = [workflow_row.workflow_name]

            # Flatten statuses horizontally
            for status in workflow_row.statuses:
                data.extend([
                    status.name,
                    status.status_type,
                    status.color,
                    status.icon,
                ])

            # Pad with empty values if this workflow has fewer statuses
            remaining = max_statuses - len(workflow_row.statuses)
            data.extend([""] * (remaining * 4))

            self._write_data_row(sheet, row_idx, data)

        # Set column widths
        widths = [35]  # Workflow name
        for _ in range(max_statuses):
            widths.extend([20, 15, 12, 12])  # Name, Type, Color, Icon
        self._set_column_widths(sheet, widths)

        # Freeze header row
        sheet.freeze_panes = "A2"

        logger.debug(
            f"Generated {self.SHEET_JOB_CUSTOM_STATUS} sheet "
            f"with {len(template.job_custom_statuses)} workflow rows"
        )

    def _generate_action_buttons_sheet(
        self,
        workbook: Workbook,
        template: ExcelImportTemplate,
    ) -> None:
        """Generate the Action Buttons tab with FieldPulse column names.

        Args:
            workbook: The workbook to add the sheet to
            template: The template containing action button data
        """
        sheet = workbook.create_sheet(self.SHEET_ACTION_BUTTONS)

        # Write header row
        self._write_header_row(sheet, self.HEADERS_ACTION_BUTTONS)

        # Write data rows
        for row_idx, button_row in enumerate(template.action_buttons, start=2):
            data = [
                button_row.workflow_name,
                button_row.action_flow_name,
                button_row.job_status_name,
                button_row.user_role,
                button_row.button_action,
                button_row.action_option or "",
                button_row.action_button_name,
            ]
            self._write_data_row(sheet, row_idx, data)

        # Set column widths
        self._set_column_widths(sheet, self.WIDTHS_ACTION_BUTTONS)

        # Freeze header row
        sheet.freeze_panes = "A2"

        logger.debug(
            f"Generated {self.SHEET_ACTION_BUTTONS} sheet "
            f"with {len(template.action_buttons)} rows"
        )

    def _generate_focus_view_sheet(
        self,
        workbook: Workbook,
        template: ExcelImportTemplate,
    ) -> None:
        """Generate the Focus View + Status Instruction tab with FieldPulse column names.

        Args:
            workbook: The workbook to add the sheet to
            template: The template containing focus view data
        """
        sheet = workbook.create_sheet(self.SHEET_FOCUS_VIEW)

        # Write header row
        self._write_header_row(sheet, self.HEADERS_FOCUS_VIEW)

        # Write data rows
        for row_idx, focus_row in enumerate(template.focus_view, start=2):
            data = [
                focus_row.workflow_name,
                focus_row.action_flow_name,
                focus_row.job_status_name,
                focus_row.user_role,
                focus_row.status_instructions,
                focus_row.display_action_menu,
                focus_row.ability_to_change_status,
                focus_row.focus_view_enabled,
                focus_row.focus_view_layout,
                focus_row.restrict_to_focus_view,
            ]
            self._write_data_row(sheet, row_idx, data, wrap_columns=[5, 9])

        # Set column widths
        self._set_column_widths(sheet, self.WIDTHS_FOCUS_VIEW)

        # Freeze header row
        sheet.freeze_panes = "A2"

        logger.debug(
            f"Generated {self.SHEET_FOCUS_VIEW} sheet "
            f"with {len(template.focus_view)} rows"
        )

    def _write_header_row(
        self,
        sheet: Worksheet,
        headers: list[str],
    ) -> None:
        """Write and format the header row.

        Args:
            sheet: The worksheet to write to
            headers: List of header strings
        """
        for col_idx, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col_idx, value=header)
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = self.DATA_ALIGNMENT
            cell.border = self.THIN_BORDER

    def _write_data_row(
        self,
        sheet: Worksheet,
        row_idx: int,
        data: list,
        wrap_columns: Optional[list[int]] = None,
    ) -> None:
        """Write and format a data row.

        Args:
            sheet: The worksheet to write to
            row_idx: The row number (1-indexed)
            data: List of values to write
            wrap_columns: Optional list of column indices (1-indexed) to apply text wrapping
        """
        wrap_columns = wrap_columns or []
        for col_idx, value in enumerate(data, start=1):
            cell = sheet.cell(row=row_idx, column=col_idx, value=value)
            if col_idx in wrap_columns:
                cell.alignment = self.WRAP_ALIGNMENT
            else:
                cell.alignment = self.DATA_ALIGNMENT
            cell.border = self.THIN_BORDER

    def _set_column_widths(
        self,
        sheet: Worksheet,
        widths: list[int],
    ) -> None:
        """Set column widths for a sheet.

        Args:
            sheet: The worksheet to modify
            widths: List of column widths
        """
        for col_idx, width in enumerate(widths, start=1):
            col_letter = self._get_column_letter(col_idx)
            sheet.column_dimensions[col_letter].width = width

    def _get_column_letter(self, col_idx: int) -> str:
        """Convert column index to Excel column letter.

        Args:
            col_idx: 1-indexed column number

        Returns:
            Column letter (A, B, ..., Z, AA, AB, etc.)
        """
        result = ""
        while col_idx > 0:
            col_idx, remainder = divmod(col_idx - 1, 26)
            result = chr(ord("A") + remainder) + result
        return result
