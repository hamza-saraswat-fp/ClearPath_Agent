"""Unit Tests for the Ingestion Pipeline.

Tests cover:
- ProductionDataParser: Excel parsing
- PatternExtractor: Pattern discovery
- EmbeddingGenerator: Rich text generation
- DataEnricher: Data enrichment
- DataValidator: Validation and deduplication
"""

import json
import tempfile
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import Workbook

from clearpath_agent.models.enums import StatusCategory, UserRole
from clearpath_agent.services.data_enricher import DataEnricher
from clearpath_agent.services.data_validator import (
    DataValidator,
    ValidationSeverity,
)
from clearpath_agent.services.embedding_generator import EmbeddingGenerator
from clearpath_agent.services.pattern_extractor import (
    FrequencyStats,
    PatternExtractor,
)
from clearpath_agent.services.prod_data_parser import (
    ActionButtonConfig,
    FocusViewConfig,
    ParsedWorkflow,
    ProductionDataParser,
    StatusConfig,
)


# ============================================================================
# Fixtures
# ============================================================================


@pytest.fixture
def sample_excel_file():
    """Create a sample Excel file for testing."""
    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = Path(tmpdir) / "test_workflow.xlsx"
        wb = Workbook()

        # Create Job Custom Status sheet
        ws_status = wb.active
        ws_status.title = "Job Custom Status"
        ws_status.append(
            [
                "Status Action Flow Name",
                "Status Name",
                "Status Category",
                "Status Color",
                "Sequence",
                "Is Active",
            ]
        )
        ws_status.append(["Test Flow", "New", "Pending", "#FF0000", 1, True])
        ws_status.append(["Test Flow", "In Progress", "In Progress", "#00FF00", 2, True])
        ws_status.append(["Test Flow", "Complete", "Completed", "#0000FF", 3, True])

        # Create Action Buttons sheet
        ws_actions = wb.create_sheet("Action Buttons")
        ws_actions.append(
            [
                "Status Action Flow Name",
                "Status Name",
                "User Role",
                "Action Type",
                "Button Label",
                "Button Order",
            ]
        )
        ws_actions.append(
            ["Test Flow", "New", "Service Agent", "Clock In", "Start Work", 1]
        )
        ws_actions.append(
            ["Test Flow", "New", "Service Agent", "Take Photo", "Before Photo", 2]
        )
        ws_actions.append(
            ["Test Flow", "In Progress", "Service Agent", "Fill Form", "Inspection", 1]
        )

        # Create Focus View sheet
        ws_focus = wb.create_sheet("Focus View")
        ws_focus.append(
            [
                "Status Action Flow Name",
                "Status Name",
                "User Role",
                "Widgets",
                "Status Instructions",
            ]
        )
        ws_focus.append(
            [
                "Test Flow",
                "New",
                "Service Agent",
                "Job Title; Customer Contact; Action Buttons",
                "Review job details",
            ]
        )
        ws_focus.append(
            [
                "Test Flow",
                "In Progress",
                "Service Agent",
                "Job Title; Forms; Files/Photos",
                "Complete the work",
            ]
        )

        wb.save(filepath)
        yield filepath


@pytest.fixture
def sample_seed_data():
    """Sample seed data for testing."""
    return {
        "actions": [
            {
                "id": "action:clock_in",
                "name": "Clock In",
                "description": "Start time tracking",
                "requires_template": False,
                "template_type": None,
                "use_cases": ["Time tracking"],
                "common_phrases": ["clock in", "start clock"],
            },
            {
                "id": "action:take_photo",
                "name": "Take Photo",
                "description": "Take a photo",
                "requires_template": False,
                "template_type": None,
                "use_cases": ["Documentation"],
                "common_phrases": ["take photo", "snap picture"],
            },
        ],
        "widgets": [
            {
                "id": "widget:job_title",
                "name": "Job Title",
                "description": "Displays job title",
                "category": "job_info",
                "typical_position": "top",
                "use_cases": ["Job identification"],
            },
            {
                "id": "widget:files_photos",
                "name": "Files/Photos",
                "description": "Shows files and photos",
                "category": "documentation",
                "typical_position": "middle",
                "use_cases": ["Photo viewing"],
            },
        ],
        "status_patterns": [
            {
                "id": "pattern:new",
                "name": "New",
                "description": "Initial status",
                "typical_status_names": ["New", "Pending"],
                "typical_actions": ["action:clock_in"],
                "typical_widgets": ["widget:job_title"],
            }
        ],
        "relationships": [
            {
                "source": "action:take_photo",
                "target": "widget:files_photos",
                "relation": "SUGGESTS_WIDGET",
                "weight": 0.9,
                "reason": "Taking photos implies viewing photos",
            }
        ],
    }


@pytest.fixture
def sample_parsed_workflow():
    """Sample ParsedWorkflow for testing."""
    return ParsedWorkflow(
        source_file="test.xlsx",
        flow_name="Test Flow",
        statuses=[
            StatusConfig(
                name="New",
                category=StatusCategory.PENDING,
                sequence=1,
                flow_name="Test Flow",
            ),
            StatusConfig(
                name="In Progress",
                category=StatusCategory.IN_PROGRESS,
                sequence=2,
                flow_name="Test Flow",
            ),
        ],
        action_buttons=[
            ActionButtonConfig(
                action_type="Clock In",
                label="Start Work",
                status_name="New",
                flow_name="Test Flow",
                user_role=UserRole.SERVICE_AGENT,
                order=1,
            ),
            ActionButtonConfig(
                action_type="Take Photo",
                label="Before Photo",
                status_name="New",
                flow_name="Test Flow",
                user_role=UserRole.SERVICE_AGENT,
                order=2,
            ),
        ],
        focus_views=[
            FocusViewConfig(
                status_name="New",
                flow_name="Test Flow",
                user_role=UserRole.SERVICE_AGENT,
                widgets=["Job Title", "Files/Photos"],
            )
        ],
    )


# ============================================================================
# ProductionDataParser Tests
# ============================================================================


class TestProductionDataParser:
    """Tests for ProductionDataParser."""

    def test_parse_valid_excel(self, sample_excel_file):
        """Parse valid Excel file returns correct ParsedWorkflow."""
        parser = ProductionDataParser()
        result = parser.parse_excel_file(sample_excel_file)

        assert isinstance(result, ParsedWorkflow)
        assert result.source_file == "test_workflow.xlsx"
        assert len(result.statuses) == 3
        assert len(result.action_buttons) == 3
        assert len(result.focus_views) == 2

    def test_parse_excel_status_names(self, sample_excel_file):
        """Parsed statuses have correct names."""
        parser = ProductionDataParser()
        result = parser.parse_excel_file(sample_excel_file)

        status_names = [s.name for s in result.statuses]
        assert "New" in status_names
        assert "In Progress" in status_names
        assert "Complete" in status_names

    def test_parse_excel_actions(self, sample_excel_file):
        """Parsed action buttons have correct data."""
        parser = ProductionDataParser()
        result = parser.parse_excel_file(sample_excel_file)

        clock_in = next(
            (a for a in result.action_buttons if a.action_type == "Clock In"), None
        )
        assert clock_in is not None
        assert clock_in.label == "Start Work"
        assert clock_in.status_name == "New"

    def test_parse_excel_widgets(self, sample_excel_file):
        """Parsed focus views have correct widgets."""
        parser = ProductionDataParser()
        result = parser.parse_excel_file(sample_excel_file)

        new_fv = next(
            (fv for fv in result.focus_views if fv.status_name == "New"), None
        )
        assert new_fv is not None
        assert "Job Title" in new_fv.widgets
        assert "Customer Contact" in new_fv.widgets

    def test_parse_nonexistent_file(self):
        """Parsing non-existent file raises FileNotFoundError."""
        parser = ProductionDataParser()
        with pytest.raises(FileNotFoundError):
            parser.parse_excel_file(Path("/nonexistent/file.xlsx"))

    def test_parse_non_excel_file(self, tmp_path):
        """Parsing non-Excel file raises ValueError."""
        text_file = tmp_path / "test.txt"
        text_file.write_text("not an excel file")

        parser = ProductionDataParser()
        with pytest.raises(ValueError):
            parser.parse_excel_file(text_file)


# ============================================================================
# PatternExtractor Tests
# ============================================================================


class TestPatternExtractor:
    """Tests for PatternExtractor."""

    def test_extract_status_flows(self, sample_parsed_workflow):
        """Extract status flows identifies sequences."""
        extractor = PatternExtractor(min_occurrences=1)
        workflows = [sample_parsed_workflow]

        flows = extractor.extract_status_flows(workflows)

        # Should find at least one flow
        assert len(flows) >= 1
        # Flow should have New -> In Progress sequence
        flow_strings = [f.as_string for f in flows]
        assert any("New" in s and "In Progress" in s for s in flow_strings)

    def test_extract_cooccurrence(self, sample_parsed_workflow):
        """Extract co-occurrences finds action-widget pairs."""
        extractor = PatternExtractor(min_occurrences=1, min_confidence=0.0)
        workflows = [sample_parsed_workflow]

        patterns = extractor.extract_action_widget_cooccurrence(workflows)

        # Should find Clock In / Take Photo paired with Job Title / Files/Photos
        action_names = [p.source_name for p in patterns]
        assert "Clock In" in action_names or "Take Photo" in action_names

    def test_calculate_frequencies(self, sample_parsed_workflow):
        """Calculate frequencies counts correctly."""
        extractor = PatternExtractor()
        workflows = [sample_parsed_workflow]

        stats = extractor.calculate_frequencies(workflows)

        assert stats.total_workflows == 1
        assert stats.action_frequencies["Clock In"] == 1
        assert stats.action_frequencies["Take Photo"] == 1
        assert "Job Title" in stats.widget_frequencies

    def test_empty_workflows(self):
        """Handle empty workflow list gracefully."""
        extractor = PatternExtractor()
        result = extractor.extract_all_patterns([])

        assert result["status_flows"] == []
        assert result["action_widget_cooccurrence"] == []
        assert result["frequencies"].total_workflows == 0


# ============================================================================
# EmbeddingGenerator Tests
# ============================================================================


class TestEmbeddingGenerator:
    """Tests for EmbeddingGenerator."""

    def test_create_action_rich_text(self, sample_seed_data):
        """Rich text for action includes key information."""
        generator = EmbeddingGenerator()
        action = sample_seed_data["actions"][0]

        rich_text = generator._create_action_rich_text(action, None)

        assert "Clock In" in rich_text
        assert "Start time tracking" in rich_text
        assert "Time tracking" in rich_text  # use case
        assert "clock in" in rich_text  # common phrase

    def test_create_widget_rich_text(self, sample_seed_data):
        """Rich text for widget includes key information."""
        generator = EmbeddingGenerator()
        widget = sample_seed_data["widgets"][0]

        rich_text = generator._create_widget_rich_text(widget, None)

        assert "Job Title" in rich_text
        assert "job_info" in rich_text  # category
        assert "top" in rich_text  # position

    def test_create_rich_text_with_patterns(self, sample_seed_data):
        """Rich text includes production frequency when available."""
        generator = EmbeddingGenerator()
        action = sample_seed_data["actions"][0]

        # Create mock patterns
        patterns = FrequencyStats()
        patterns.action_frequencies["Clock In"] = 10
        patterns.total_workflows = 5

        rich_text = generator._create_action_rich_text(action, patterns)

        assert "10" in rich_text  # frequency count

    @patch.object(EmbeddingGenerator, "_call_embedding_api")
    def test_generate_action_embedding(self, mock_api, sample_seed_data):
        """Generate embedding creates ActionEmbedding object."""
        mock_api.return_value = [0.1] * 1536

        generator = EmbeddingGenerator()
        action = sample_seed_data["actions"][0]

        result = generator.generate_action_embedding(action, None)

        assert result.id == "action:clock_in"
        assert result.name == "Clock In"
        assert len(result.embedding) == 1536
        assert result.rich_text != ""


# ============================================================================
# DataEnricher Tests
# ============================================================================


class TestDataEnricher:
    """Tests for DataEnricher."""

    def test_enrich_actions_with_frequency(self, sample_seed_data):
        """Enrich actions adds production frequency."""
        enricher = DataEnricher()

        # Create patterns with matching action
        patterns = FrequencyStats()
        patterns.action_frequencies["Clock In"] = 15
        patterns.total_workflows = 3

        enriched = enricher.enrich_actions(sample_seed_data["actions"], patterns)

        clock_in = next(a for a in enriched if a["name"] == "Clock In")
        assert clock_in["production_frequency"] == 15
        assert clock_in["production_workflows"] == 3

    def test_enrich_actions_fuzzy_match(self, sample_seed_data):
        """Enrich actions uses fuzzy matching for similar names."""
        enricher = DataEnricher(fuzzy_threshold=0.8)

        # Pattern with slightly different name
        patterns = FrequencyStats()
        patterns.action_frequencies["Clock-In"] = 10  # Hyphenated variant
        patterns.total_workflows = 2

        enriched = enricher.enrich_actions(sample_seed_data["actions"], patterns)

        clock_in = next(a for a in enriched if a["name"] == "Clock In")
        # Should find fuzzy match
        assert clock_in["production_frequency"] >= 0

    def test_discover_new_entities(self, sample_seed_data, sample_parsed_workflow):
        """Discover new entities finds unknown actions/widgets."""
        enricher = DataEnricher()

        # Add an action not in seed data
        sample_parsed_workflow.action_buttons.append(
            ActionButtonConfig(
                action_type="New Action Type",
                label="Do Something New",
                status_name="New",
                flow_name="Test",
            )
        )

        discovered = enricher.discover_new_entities(
            [sample_parsed_workflow], sample_seed_data
        )

        # Should discover the new action
        new_action_names = [a["name"] for a in discovered["new_actions"]]
        assert "New Action Type" in new_action_names


# ============================================================================
# DataValidator Tests
# ============================================================================


class TestDataValidator:
    """Tests for DataValidator."""

    def test_validate_valid_action(self, sample_seed_data):
        """Valid action passes validation."""
        validator = DataValidator()
        action = sample_seed_data["actions"][0]

        result = validator.validate_action(action)

        assert result.is_valid
        assert len(result.errors) == 0

    def test_validate_action_missing_field(self):
        """Action missing required field fails validation."""
        validator = DataValidator()
        action = {"id": "action:test"}  # Missing name and description

        result = validator.validate_action(action)

        assert not result.is_valid
        error_fields = [e.field for e in result.errors]
        assert "name" in error_fields
        assert "description" in error_fields

    def test_validate_action_invalid_id(self):
        """Action with invalid ID format generates warning."""
        validator = DataValidator()
        action = {
            "id": "invalid_id",  # Should start with action:
            "name": "Test",
            "description": "Test description",
        }

        result = validator.validate_action(action)

        # Should have warning about ID format
        warning_fields = [w.field for w in result.warnings]
        assert "id" in warning_fields

    def test_validate_action_requires_template(self):
        """Action with requires_template but no template_type fails."""
        validator = DataValidator()
        action = {
            "id": "action:test",
            "name": "Test",
            "description": "Test",
            "requires_template": True,
            # Missing template_type
        }

        result = validator.validate_action(action)

        assert not result.is_valid
        error_fields = [e.field for e in result.errors]
        assert "template_type" in error_fields

    def test_validate_valid_widget(self, sample_seed_data):
        """Valid widget passes validation."""
        validator = DataValidator()
        widget = sample_seed_data["widgets"][0]

        result = validator.validate_widget(widget)

        assert result.is_valid

    def test_validate_widget_missing_category(self):
        """Widget missing category fails validation."""
        validator = DataValidator()
        widget = {
            "id": "widget:test",
            "name": "Test",
            "description": "Test",
            # Missing category
        }

        result = validator.validate_widget(widget)

        assert not result.is_valid

    def test_validate_relationship_invalid_source(self, sample_seed_data):
        """Relationship with invalid source fails validation."""
        validator = DataValidator()
        known_ids = {
            a["id"] for a in sample_seed_data["actions"]
        } | {w["id"] for w in sample_seed_data["widgets"]}

        rel = {
            "source": "action:nonexistent",
            "target": "widget:job_title",
            "relation": "SUGGESTS_WIDGET",
            "weight": 0.5,
        }

        result = validator.validate_relationship(rel, known_ids)

        assert not result.is_valid
        error_fields = [e.field for e in result.errors]
        assert "source" in error_fields

    def test_validate_relationship_invalid_weight(self, sample_seed_data):
        """Relationship with weight out of bounds fails validation."""
        validator = DataValidator()
        known_ids = {"action:clock_in", "widget:job_title"}

        rel = {
            "source": "action:clock_in",
            "target": "widget:job_title",
            "relation": "SUGGESTS_WIDGET",
            "weight": 1.5,  # Invalid: should be 0-1
        }

        result = validator.validate_relationship(rel, known_ids)

        assert not result.is_valid

    def test_deduplicate_exact_match(self):
        """Deduplication removes exact ID matches."""
        validator = DataValidator()
        actions = [
            {"id": "action:test", "name": "Test", "description": "Original"},
            {"id": "action:test", "name": "Test", "description": "Duplicate"},
        ]

        deduped, result = validator.deduplicate_actions(actions)

        assert len(deduped) == 1
        assert result.exact_matches == 1

    def test_deduplicate_fuzzy_match(self):
        """Deduplication flags fuzzy matches for review."""
        validator = DataValidator(fuzzy_threshold=0.9)
        actions = [
            {"id": "action:clock_in", "name": "Clock In", "description": "Start"},
            {"id": "action:clockin", "name": "ClockIn", "description": "Start work"},
        ]

        deduped, result = validator.deduplicate_actions(actions)

        # Should flag the fuzzy match
        assert result.fuzzy_matches == 1
        assert len(result.flagged_for_review) == 1

    def test_generate_validation_report(self, sample_seed_data):
        """Generate validation report produces readable output."""
        validator = DataValidator()
        results = validator.validate_all(
            sample_seed_data["actions"],
            sample_seed_data["widgets"],
            sample_seed_data["status_patterns"],
            sample_seed_data["relationships"],
        )

        report = validator.generate_validation_report(results)

        assert "Validation Report" in report
        assert "Total entities validated" in report


# ============================================================================
# Integration Tests
# ============================================================================


class TestPipelineIntegration:
    """Integration tests for the full pipeline."""

    def test_parse_extract_enrich_flow(
        self, sample_excel_file, sample_seed_data
    ):
        """Test full flow from parsing to enrichment."""
        # Parse
        parser = ProductionDataParser()
        workflow = parser.parse_excel_file(sample_excel_file)

        # Extract patterns
        extractor = PatternExtractor(min_occurrences=1, min_confidence=0.0)
        patterns = extractor.extract_all_patterns([workflow])

        # Enrich
        enricher = DataEnricher()
        enriched_actions = enricher.enrich_actions(
            sample_seed_data["actions"], patterns["frequencies"]
        )

        # Validate
        validator = DataValidator()
        results = validator.validate_all(
            enriched_actions,
            sample_seed_data["widgets"],
            sample_seed_data["status_patterns"],
            sample_seed_data["relationships"],
        )

        # All seed data should be valid
        assert all(r.is_valid for r in results)

    def test_pipeline_with_empty_prod_data(self, sample_seed_data):
        """Pipeline handles case with no production data."""
        # Empty production data
        extractor = PatternExtractor()
        patterns = extractor.extract_all_patterns([])

        # Should still work with empty patterns
        enricher = DataEnricher()
        enriched = enricher.enrich_actions(
            sample_seed_data["actions"], patterns["frequencies"]
        )

        # Actions should have 0 frequency
        for action in enriched:
            assert action.get("production_frequency", 0) == 0
